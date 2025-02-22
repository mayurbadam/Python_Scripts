import os
import datetime
import pandas as pd
import matplotlib
matplotlib.use("Agg")  # Use non-GUI backend
import matplotlib.pyplot as plt
from flask import Flask, render_template, request, redirect, url_for
import base64
from io import BytesIO
from openpyxl import load_workbook

app = Flask(__name__)

# Excel file for local storage
excel_file = "expenses.xlsx"

def load_data():
    month_name = datetime.datetime.now().strftime("%B")
    if os.path.exists(excel_file):
        with pd.ExcelFile(excel_file, engine="openpyxl") as xls:
            if month_name in xls.sheet_names:
                return pd.read_excel(xls, sheet_name=month_name, engine="openpyxl")
    return pd.DataFrame(columns=["Date", "Name", "Rent", "Need", "Transit", "Luxury", "Lent", "Savings"])

def save_data(data):
    month_name = datetime.datetime.now().strftime("%B")
    with pd.ExcelWriter(excel_file, engine="openpyxl", mode="a") as writer:
        writer.book = load_workbook(excel_file)
        if month_name in writer.book.sheetnames:
            writer.book.remove(writer.book[month_name])
        data.to_excel(writer, sheet_name=month_name, index=False)
        writer._save()

def generate_chart(data):
    categories = ["Rent", "Need", "Transit", "Luxury", "Lent", "Savings"]
    values = [data[col].astype(float).sum() for col in categories]
    
    print(f"Chart Data: {values}")  # Debugging: Check values being passed
    
    if sum(values) == 0:
        print("No expenses to plot.")  # Debugging
        return None
    
    fig, ax = plt.subplots(figsize=(5, 3))
    ax.pie(values, labels=categories, autopct="%1.1f%%", startangle=90, 
           colors=["#ff9999", "#66b3ff", "#99ff99", "#ffcc99", "#c2c2f0", "#ffb3e6"])
    ax.set_title("Expense Distribution")
    
    buf = BytesIO()
    plt.savefig(buf, format="png", bbox_inches='tight')
    buf.seek(0)
    plt.close(fig)  # Free memory
    
    return base64.b64encode(buf.getvalue()).decode("utf-8")

@app.route('/')
def index():
    data = load_data()
    summary = {col: data[col].astype(float).sum() for col in ["Rent", "Need", "Transit", "Luxury", "Lent", "Savings"]} if not data.empty else {}
    chart = generate_chart(data)
    return render_template("index.html", data=data.to_dict(orient='records'), summary=summary, chart=chart)

@app.route('/add', methods=['POST'])
def add_expense():
    data = load_data()
    new_row = pd.DataFrame([{ 
        "Date": request.form.get("date") or datetime.datetime.now().strftime("%d-%b"),
        "Name": request.form.get("name") or " ",
        "Rent": request.form.get("rent") or "0",
        "Need": request.form.get("need") or "0",
        "Transit": request.form.get("transit") or "0",
        "Luxury": request.form.get("luxury") or "0",
        "Lent": request.form.get("lent") or "0",
        "Savings": request.form.get("savings") or "0"
    }])
    data = pd.concat([data, new_row], ignore_index=True)
    save_data(data)
    return redirect(url_for("index"))

if __name__ == "__main__":
    app.run(debug=True)

